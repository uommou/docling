# SPDX-FileCopyrightText: The Docling Contributors
# SPDX-License-Identifier: MIT

"""Tests for PowerPoint shapes filled with a picture (``a:blipFill``)."""

from io import BytesIO
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from .test_backend_pptx import get_converter


def _add_picture_filled_rectangle(slide, blip_attr: str, image_rid: str, text: str):
    """Add a rectangle whose ``p:spPr`` fills it with the picture ``image_rid``."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2)
    )
    blip_fill = etree.Element(qn("a:blipFill"))
    etree.SubElement(blip_fill, qn("a:blip")).set(qn(blip_attr), image_rid)
    etree.SubElement(etree.SubElement(blip_fill, qn("a:stretch")), qn("a:fillRect"))
    shape._element.spPr.find(qn("a:prstGeom")).addnext(blip_fill)
    shape.text_frame.text = text
    return shape


@pytest.mark.parametrize("text", ["", "Label on the picture"])
def test_pptx_picture_filled_shape(tmp_path: Path, text: str):
    """A picture-filled shape gives a picture at its position and keeps its text."""
    png = BytesIO()
    Image.new("RGB", (200, 100), (200, 30, 30)).save(png, format="PNG")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _, image_rid = slide.part.get_or_add_image_part(png)
    shape = _add_picture_filled_rectangle(slide, "r:embed", image_rid, text)

    pptx_path = tmp_path / "picture_fill.pptx"
    prs.save(pptx_path)

    doc = get_converter().convert(pptx_path).document

    assert len(doc.pictures) == 1
    picture = doc.pictures[0]
    assert picture.image is not None
    assert picture.image.size.width == 200.0
    assert picture.image.size.height == 100.0
    bbox = picture.prov[0].bbox
    assert (bbox.l, bbox.t, bbox.r, bbox.b) == (
        shape.left,
        shape.top,
        shape.left + shape.width,
        shape.top + shape.height,
    )
    assert [item.text for item in doc.texts] == ([text] if text else [])


def test_pptx_linked_picture_fill_is_skipped_with_warning(tmp_path: Path):
    """A picture fill that only links an external image is skipped with a warning."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image_rid = slide.part.relate_to(
        "https://example.com/picture.png", RT.IMAGE, is_external=True
    )
    _add_picture_filled_rectangle(slide, "r:link", image_rid, "Linked fill")

    pptx_path = tmp_path / "linked_picture_fill.pptx"
    prs.save(pptx_path)

    with pytest.warns(UserWarning, match="picture fill without an embedded image"):
        doc = get_converter().convert(pptx_path).document

    assert doc.pictures == []
    assert [item.text for item in doc.texts] == ["Linked fill"]
